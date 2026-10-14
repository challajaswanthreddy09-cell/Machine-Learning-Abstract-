"""
backend/prediction_service.py
Core inference and Excel persistence engine for Loan Default Prediction.
"""

import os
import json
import threading
from datetime import datetime
import joblib
import pandas as pd
from openpyxl import Workbook, load_workbook

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MODEL_PATH = os.path.join(BASE_DIR, "ml", "models", "loan_default_model.joblib")
METADATA_PATH = os.path.join(BASE_DIR, "ml", "model_metadata.json")
CLUSTERING_PATH = os.path.join(BASE_DIR, "ml", "clustering_summary.json")
EXCEL_PATH = os.path.join(BASE_DIR, "data", "prediction_history.xlsx")

# Excel file writing lock to ensure thread safety
excel_lock = threading.Lock()


class PredictionService:
    def __init__(self):
        self.pipeline = None
        self.metadata = {}
        self.clustering_summary = {}
        self.load_artifacts()
        self._ensure_excel_file()

    def load_artifacts(self):
        """Loads serialized scikit-learn pipeline and metadata JSON."""
        if not os.path.exists(MODEL_PATH):
            raise FileNotFoundError(f"Model file not found at {MODEL_PATH}. Please run ml/train.py first.")
        
        self.pipeline = joblib.load(MODEL_PATH)
        print(f"[PredictionService] Successfully loaded trained pipeline from: {MODEL_PATH}")

        if os.path.exists(METADATA_PATH):
            with open(METADATA_PATH, "r", encoding="utf-8") as f:
                self.metadata = json.load(f)
            print(f"[PredictionService] Loaded metadata for model: {self.metadata.get('model_name')}")

        if os.path.exists(CLUSTERING_PATH):
            with open(CLUSTERING_PATH, "r", encoding="utf-8") as f:
                self.clustering_summary = json.load(f)

    def _ensure_excel_file(self):
        """Initializes the Excel history workbook if it does not exist."""
        with excel_lock:
            if not os.path.exists(EXCEL_PATH):
                os.makedirs(os.path.dirname(EXCEL_PATH), exist_ok=True)
                wb = Workbook()
                ws = wb.active
                ws.title = "Prediction History"
                headers = [
                    "Timestamp",
                    "Loan Amount ($)",
                    "Annual Income ($)",
                    "Interest Rate (%)",
                    "DTI Ratio",
                    "Monthly Installment ($)",
                    "Employment Length (Years)",
                    "Credit Grade",
                    "Home Ownership",
                    "Loan Term",
                    "Loan Purpose",
                    "Application Type",
                    "Default Probability",
                    "Risk Level",
                    "Eligibility Decision",
                    "Model Used"
                ]
                ws.append(headers)
                wb.save(EXCEL_PATH)
                print(f"[PredictionService] Initialized Excel history log at: {EXCEL_PATH}")

    def log_prediction_to_excel(self, record: dict):
        """
        Thread-safely appends a new prediction record to prediction_history.xlsx.
        Preserves all prior history without overwriting.
        """
        with excel_lock:
            row_data = [
                record.get("timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                record.get("loan_amount"),
                record.get("annual_inc"),
                record.get("interest_rate"),
                record.get("dti"),
                record.get("installment"),
                record.get("emp_length_int"),
                record.get("grade"),
                record.get("home_ownership"),
                record.get("term"),
                record.get("purpose"),
                record.get("application_type"),
                f"{record.get('default_probability') * 100:.2f}%",
                record.get("risk_level"),
                record.get("decision"),
                record.get("model_used")
            ]
            wb = load_workbook(EXCEL_PATH)
            ws = wb.active
            ws.append(row_data)
            wb.save(EXCEL_PATH)

    def predict(self, input_data: dict) -> dict:
        """
        Processes applicant details through the saved ML pipeline,
        computes default probability, determines risk level,
        logs to Excel, and returns structured result.
        """
        # Convert dictionary to a 1-row DataFrame with expected columns
        feature_order = self.metadata.get("features", {}).get("numerical", []) + \
                        self.metadata.get("features", {}).get("categorical", [])
        
        # Fallback if metadata not present
        if not feature_order:
            feature_order = [
                "loan_amount", "emp_length_int", "annual_inc", "interest_rate", "dti", "installment",
                "grade", "home_ownership", "term", "purpose", "application_type"
            ]

        # Clean/standardize input values
        clean_input = {}
        for col in feature_order:
            val = input_data.get(col)
            if col in ["loan_amount", "emp_length_int", "annual_inc", "interest_rate", "dti", "installment"]:
                clean_input[col] = float(val)
            else:
                clean_input[col] = str(val).strip()

        df_input = pd.DataFrame([clean_input])

        # Generate prediction and probability
        probabilities = self.pipeline.predict_proba(df_input)[0]
        # Class 1 is Default (Bad Loan)
        default_prob = float(probabilities[1])

        # Risk Threshold Evaluation
        thresholds = self.metadata.get("risk_thresholds", {})
        low_thresh = thresholds.get("low_risk_threshold", 0.35)
        med_thresh = thresholds.get("medium_risk_threshold", 0.60)

        if default_prob < low_thresh:
            risk_level = "Low"
            eligible = True
            decision = "Eligible for Loan"
        elif default_prob < med_thresh:
            risk_level = "Medium"
            eligible = True
            decision = "Eligible for Loan (Conditional Review)"
        else:
            risk_level = "High"
            eligible = False
            decision = "Not Eligible for Loan"

        model_used = self.metadata.get("model_name", "XGBoost Classifier")
        timestamp_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        result = {
            "eligible": eligible,
            "decision": decision,
            "default_probability": round(default_prob, 4),
            "default_probability_percentage": f"{default_prob * 100:.2f}%",
            "risk_level": risk_level,
            "model_used": model_used,
            "timestamp": timestamp_str,
            "input_data": clean_input
        }

        # Persist to Excel history log
        log_entry = {**clean_input, **result}
        self.log_prediction_to_excel(log_entry)

        return result

    def get_history(self, limit: int = 50) -> list:
        """Reads recent prediction history records from Excel."""
        with excel_lock:
            if not os.path.exists(EXCEL_PATH):
                return []
            try:
                df_hist = pd.read_excel(EXCEL_PATH)
                if df_hist.empty:
                    return []
                # Return newest records first
                records = df_hist.tail(limit).to_dict(orient="records")
                records.reverse()
                return records
            except Exception as e:
                print(f"[PredictionService] Error reading history: {e}")
                return []


# Global singleton instance
service = PredictionService()
